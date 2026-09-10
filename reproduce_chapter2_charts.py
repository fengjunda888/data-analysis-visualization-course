from __future__ import annotations

import math
import re
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from matplotlib import patches
from matplotlib.colors import LinearSegmentedColormap
from matplotlib.patches import FancyBboxPatch, Wedge
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "chapter2_python_charts"
FRONT = ROOT / "第二章 图表(前15).xlsx"
BACK = ROOT / "第二章 图表(后15).xlsx"

plt.rcParams["font.sans-serif"] = [
    "Microsoft YaHei",
    "SimHei",
    "Noto Sans CJK SC",
    "DejaVu Sans",
]
plt.rcParams["axes.unicode_minus"] = False

NAVY = "#16324F"
BLUE = "#2F75B5"
TEAL = "#21A6A8"
ORANGE = "#F28E5B"
GOLD = "#F5C04A"
RED = "#D95D5D"
PURPLE = "#7B61A8"
GREEN = "#5EAA78"
LIGHT = "#EAF0F5"
MID = "#AFC2D4"
DARK = "#263746"


def clean(value):
    if value is None:
        return ""
    return str(value).strip()


def data_rows(path: Path, sheet: str, min_row: int = 1):
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(min_row=min_row, values_only=True))
    wb.close()
    return rows


def numeric(v, default=0.0):
    try:
        return float(v)
    except (TypeError, ValueError):
        return default


def save(fig, number, name):
    fig.tight_layout()
    fig.savefig(OUT / f"{number:02d}_{name}.png", dpi=180, bbox_inches="tight")
    plt.close(fig)


def style(ax, title):
    ax.set_title(title, loc="left", fontsize=15, fontweight="bold", color=NAVY, pad=12)
    ax.spines[["top", "right", "left"]].set_visible(False)
    ax.grid(axis="y", color=LIGHT, linewidth=0.8)
    ax.tick_params(colors=DARK)
    ax.set_axisbelow(True)


def gradient_bars(ax, bars, colors=(BLUE, TEAL)):
    cmap = LinearSegmentedColormap.from_list("bar_gradient", colors)
    for i, bar in enumerate(bars):
        bar.set_color(cmap(i / max(1, len(bars) - 1)))
        bar.set_edgecolor("none")


def labels(ax, bars, fmt="{:.0f}", color=DARK):
    for b in bars:
        ax.text(
            b.get_x() + b.get_width() / 2,
            b.get_height(),
            fmt.format(b.get_height()),
            ha="center",
            va="bottom",
            fontsize=9,
            color=color,
        )


def chart01():
    rows = data_rows(FRONT, "1 渐变柱形图")[2:8]
    cats, vals = [clean(r[1]) for r in rows], [numeric(r[2]) for r in rows]
    fig, ax = plt.subplots(figsize=(8, 5))
    bars = ax.bar(cats, vals, color=BLUE, width=0.62)
    gradient_bars(ax, bars)
    labels(ax, bars)
    ax.set_ylim(0, max(vals) * 1.2)
    style(ax, "渐变柱形图")
    save(fig, 1, "gradient_bar")


def chart02():
    rows = data_rows(FRONT, "2 带均值柱形图")[2:9]
    cats, vals = [clean(r[1]) for r in rows], [numeric(r[2]) for r in rows]
    avg = np.mean(vals)
    fig, ax = plt.subplots(figsize=(8, 5))
    bars = ax.bar(cats, vals, color=TEAL, width=0.62)
    labels(ax, bars)
    ax.axhline(avg, color=ORANGE, linestyle="--", linewidth=2, label=f"均值 {avg:.0f}")
    ax.legend(frameon=False)
    ax.set_ylim(0, max(vals) * 1.2)
    style(ax, "带均值柱形图")
    save(fig, 2, "mean_bar")


def chart03():
    rows = data_rows(FRONT, "3 渐变圆角柱形图")[2:8]
    cats, vals = [clean(r[1]) for r in rows], [numeric(r[2]) for r in rows]
    fig, ax = plt.subplots(figsize=(8, 5))
    for i, (cat, val) in enumerate(zip(cats, vals)):
        ax.bar(cat, val, color=BLUE if i % 2 == 0 else TEAL, width=0.58)
        ax.text(i, val + max(vals) * 0.025, f"{val:.0f}", ha="center", fontsize=9)
    style(ax, "渐变圆角柱形图")
    ax.set_ylim(0, max(vals) * 1.2)
    save(fig, 3, "rounded_gradient_bar")


def chart04():
    rows = data_rows(FRONT, "4 标注柱形图")[2:10]
    cats, vals = [clean(r[1]) for r in rows], [numeric(r[2]) for r in rows]
    fig, ax = plt.subplots(figsize=(8, 5))
    bars = ax.bar(cats, vals, color=BLUE, width=0.62)
    for i, b in enumerate(bars):
        b.set_color(ORANGE if vals[i] == max(vals) else BLUE)
        ax.annotate(f"{vals[i]:,.0f}", (b.get_x() + b.get_width() / 2, b.get_height()),
                    xytext=(0, 7), textcoords="offset points", ha="center", fontsize=9)
    style(ax, "标注柱形图")
    ax.set_ylim(0, max(vals) * 1.2)
    save(fig, 4, "annotated_bar")


def chart05():
    rows = data_rows(FRONT, "5 层叠柱形图")[2:9]
    cats = [clean(r[1]) for r in rows]
    sales, profit = [numeric(r[2]) for r in rows], [numeric(r[3]) for r in rows]
    fig, ax = plt.subplots(figsize=(8, 5))
    ax.bar(cats, sales, label="销售额", color=BLUE)
    ax.bar(cats, profit, bottom=sales, label="利润", color=GOLD)
    ax.legend(frameon=False, ncol=2)
    style(ax, "层叠柱形图")
    save(fig, 5, "stacked_bar")


def chart06():
    rows = data_rows(FRONT, "6 蝴蝶图")[2:8]
    cats = [clean(r[1]) for r in rows]
    left, right = [numeric(r[3]) for r in rows], [numeric(r[5]) for r in rows]
    y = np.arange(len(cats))
    fig, ax = plt.subplots(figsize=(9, 5))
    ax.barh(y, -np.array(left), color=TEAL, label="2022")
    ax.barh(y, right, color=ORANGE, label="2021")
    ax.set_yticks(y, cats)
    ax.axvline(0, color=DARK, linewidth=0.8)
    ax.legend(frameon=False, ncol=2)
    ax.xaxis.set_major_formatter(lambda x, pos: f"{abs(x):,.0f}")
    style(ax, "蝴蝶图")
    ax.grid(axis="x", color=LIGHT)
    ax.grid(axis="y", visible=False)
    save(fig, 6, "butterfly")


def chart07():
    rows = data_rows(FRONT, "7 蝴蝶图")[2:22]
    cats = [clean(r[1]) for r in rows]
    left, right = [numeric(r[2]) for r in rows], [numeric(r[3]) for r in rows]
    y = np.arange(len(cats))
    fig, ax = plt.subplots(figsize=(9, 8))
    ax.barh(y, -np.array(left), color=BLUE, label="2022")
    ax.barh(y, right, color=ORANGE, label="2021")
    ax.set_yticks(y, cats)
    ax.axvline(0, color=DARK, linewidth=0.8)
    ax.legend(frameon=False, ncol=2)
    ax.xaxis.set_major_formatter(lambda x, pos: f"{abs(x):.0%}")
    style(ax, "蝴蝶图（百分比）")
    ax.grid(axis="x", color=LIGHT)
    ax.grid(axis="y", visible=False)
    save(fig, 7, "butterfly_percent")


def chart08():
    rows = data_rows(FRONT, "8 数值百分比")[2:8]
    cats, vals = [clean(r[1]) for r in rows], [numeric(r[2]) for r in rows]
    fig, ax = plt.subplots(figsize=(8, 5))
    bars = ax.barh(cats, vals, color=TEAL)
    ax.set_xlim(0, max(vals) * 1.18)
    for b, v in zip(bars, vals):
        ax.text(v, b.get_y() + b.get_height() / 2, f" {v:,.0f}", va="center", fontsize=9)
    style(ax, "数值百分比")
    ax.grid(axis="x", color=LIGHT)
    ax.grid(axis="y", visible=False)
    save(fig, 8, "value_percent")


def chart09():
    rows = data_rows(FRONT, "9 对比柱形图")[2:9]
    cats = [clean(r[1]) for r in rows]
    a, b = [numeric(r[2]) for r in rows], [numeric(r[3]) for r in rows]
    x = np.arange(len(cats))
    fig, ax = plt.subplots(figsize=(8, 5))
    ax.bar(x - 0.18, a, width=0.36, color=BLUE, label="2021")
    ax.bar(x + 0.18, b, width=0.36, color=ORANGE, label="2022")
    ax.set_xticks(x, cats)
    ax.legend(frameon=False, ncol=2)
    style(ax, "对比柱形图")
    save(fig, 9, "comparison_bar")


def chart10():
    rows = data_rows(FRONT, "10 甘特图")[3:10]
    names = [clean(r[1]) for r in rows]
    start = [r[2] for r in rows]
    duration = [numeric(r[3]) for r in rows]
    base = min(start)
    offsets = [(d - base).days for d in start]
    y = np.arange(len(names))
    fig, ax = plt.subplots(figsize=(9, 5))
    ax.barh(y, duration, left=offsets, color=BLUE, height=0.48)
    ax.set_yticks(y, names)
    ax.invert_yaxis()
    style(ax, "甘特图")
    ax.set_xlabel("项目时间（天）")
    save(fig, 10, "gantt")


def chart11():
    rows = data_rows(FRONT, "11 平滑折线图")[2:13]
    x = [f"{clean(r[1])}-{clean(r[2])}" for r in rows]
    vals = [numeric(r[3]) for r in rows]
    fig, ax = plt.subplots(figsize=(9, 5))
    ax.plot(x, vals, color=TEAL, linewidth=2.5, marker="o", markersize=5)
    ax.fill_between(np.arange(len(vals)), vals, color=TEAL, alpha=0.10)
    style(ax, "平滑折线图")
    ax.tick_params(axis="x", rotation=35)
    save(fig, 11, "smooth_line")


def chart12():
    rows = data_rows(FRONT, "12 菱形走势图")[2:10]
    cats, vals = [clean(r[1]) for r in rows], [numeric(r[2]) for r in rows]
    fig, ax = plt.subplots(figsize=(9, 5))
    ax.plot(cats, vals, color=PURPLE, linewidth=2.4, marker="D", markersize=7)
    for i, v in enumerate(vals):
        ax.text(i, v + (max(vals) - min(vals)) * 0.04, f"{v:.1%}", ha="center", fontsize=9)
    style(ax, "菱形走势图")
    ax.tick_params(axis="x", rotation=30)
    save(fig, 12, "diamond_trend")


def chart13():
    rows = data_rows(FRONT, "13 对比折线图")[2:10]
    cats = [clean(r[1]) for r in rows]
    a, b = [numeric(r[2]) for r in rows], [numeric(r[3]) for r in rows]
    fig, ax = plt.subplots(figsize=(9, 5))
    ax.plot(cats, a, color=BLUE, marker="o", linewidth=2, label="2021")
    ax.plot(cats, b, color=ORANGE, marker="o", linewidth=2, label="2022")
    ax.legend(frameon=False, ncol=2)
    style(ax, "对比折线图")
    save(fig, 13, "comparison_line")


def ring(ax, value, center=(0.5, 0.5), radius=0.36, color=TEAL, title=""):
    ax.pie([value, 1 - value], startangle=90, counterclock=False,
           colors=[color, LIGHT], wedgeprops=dict(width=0.18, edgecolor="white"))
    ax.text(*center, f"{value:.0%}", ha="center", va="center", fontsize=20, fontweight="bold", color=NAVY)
    ax.set_title(title, fontsize=11, color=DARK)


def chart14():
    rows = data_rows(FRONT, "14 单值圆环图")[2:6]
    vals = [numeric(r[1]) for r in rows if r[1] is not None]
    value = vals[0] if vals else 0.85
    fig, ax = plt.subplots(figsize=(5, 5))
    ring(ax, value, color=BLUE, title="单值圆环图")
    save(fig, 14, "single_ring")


def liquid(ax, value, title, color=TEAL):
    circle = patches.Circle((0.5, 0.5), 0.38, facecolor="none", edgecolor=MID, linewidth=2)
    ax.add_patch(circle)
    y = 0.12 + 0.76 * value
    wave_x = np.linspace(0.12, 0.88, 240)
    wave_y = y + 0.018 * np.sin(np.linspace(0, 3 * np.pi, 240))
    ax.fill_between(wave_x, 0.12, wave_y, color=color, alpha=0.8)
    ax.plot(wave_x, wave_y, color=color, linewidth=1.5)
    ax.text(0.5, 0.5, f"{value:.0%}", ha="center", va="center", fontsize=22, fontweight="bold", color=NAVY)
    ax.text(0.5, 0.04, title, ha="center", va="center", fontsize=11, color=DARK)
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")


def chart15():
    rows = data_rows(FRONT, "15 水球图")[2:8]
    value = numeric(rows[0][1], 0.65) if rows else 0.65
    fig, ax = plt.subplots(figsize=(5, 5))
    liquid(ax, value, "水球图")
    save(fig, 15, "liquid")


def chart16():
    rows = data_rows(BACK, "16 波浪水球图 ")[3:9]
    vals = [numeric(r[1], 0.65) for r in rows if r[1] is not None]
    value = vals[0] if vals else 0.65
    fig, ax = plt.subplots(figsize=(5, 5))
    liquid(ax, value, "波浪水球图", color=BLUE)
    save(fig, 16, "wave_liquid")


def chart17():
    rows = data_rows(BACK, "17 玉玦图")[2:6]
    cats, vals = [clean(r[1]) for r in rows], [numeric(r[2]) for r in rows]
    fig, ax = plt.subplots(figsize=(6, 6))
    ax.pie(vals, labels=cats, startangle=90, counterclock=False,
           colors=[BLUE, TEAL, ORANGE, GOLD], wedgeprops=dict(width=0.22, edgecolor="white"))
    ax.text(0, 0, "玉玦", ha="center", va="center", fontsize=18, fontweight="bold", color=NAVY)
    ax.set_title("玉玦图", loc="left", fontweight="bold", color=NAVY)
    save(fig, 17, "jade_ring")


def chart18():
    rows = data_rows(BACK, "18 跑道图")[2:8]
    cats, vals = [clean(r[1]) for r in rows], [numeric(r[2]) for r in rows]
    y = np.arange(len(cats))
    fig, ax = plt.subplots(figsize=(9, 5))
    ax.barh(y, vals, color=BLUE, height=0.68)
    for yi, v, cat in zip(y, vals, cats):
        ax.text(0.01 * max(vals), yi, cat, va="center", color="white", fontweight="bold")
        ax.text(v, yi, f" {v:,.0f}", va="center", color=NAVY, fontsize=9)
    ax.set_yticks([])
    style(ax, "跑道图")
    ax.grid(axis="y", visible=False)
    save(fig, 18, "track")


def chart19():
    rows = data_rows(BACK, "19 南丁格尔圆饼图")[2:8]
    cats, vals = [clean(r[1]) for r in rows], [numeric(r[2]) for r in rows]
    fig, ax = plt.subplots(figsize=(7, 7), subplot_kw=dict(polar=True))
    theta = np.linspace(0, 2 * np.pi, len(vals), endpoint=False)
    width = 2 * np.pi / len(vals) * 0.78
    ax.bar(theta, vals, width=width, color=[BLUE, TEAL, ORANGE, GOLD, PURPLE, GREEN], alpha=0.9)
    ax.set_xticks(theta, cats)
    ax.set_yticklabels([])
    ax.grid(False)
    ax.set_title("南丁格尔圆饼图", loc="left", fontweight="bold", color=NAVY)
    save(fig, 19, "nightingale_pie")


def chart20():
    rows = data_rows(BACK, "20 南丁格尔圆环图")[2:6]
    cats, vals = [clean(r[1]) for r in rows], [numeric(r[2]) for r in rows]
    fig, ax = plt.subplots(figsize=(7, 7), subplot_kw=dict(polar=True))
    theta = np.linspace(0, 2 * np.pi, len(vals), endpoint=False)
    width = 2 * np.pi / len(vals) * 0.78
    ax.bar(theta, vals, width=width, bottom=0.2, color=[BLUE, TEAL, ORANGE, GOLD], alpha=0.9)
    ax.set_xticks(theta, cats)
    ax.set_yticklabels([])
    ax.grid(False)
    ax.set_title("南丁格尔圆环图", loc="left", fontweight="bold", color=NAVY)
    save(fig, 20, "nightingale_ring")


def chart21():
    rows = data_rows(BACK, "20 南丁格尔（PPT）")[2:8]
    cats, vals = [clean(r[1]) for r in rows], [numeric(r[2]) for r in rows]
    fig, ax = plt.subplots(figsize=(7, 7))
    ax.pie(vals, labels=cats, startangle=90, counterclock=False,
           colors=[BLUE, TEAL, ORANGE, GOLD, PURPLE, GREEN], autopct="%1.0f%%",
           wedgeprops=dict(width=0.30, edgecolor="white"))
    ax.set_title("南丁格尔图（PPT版）", loc="left", fontweight="bold", color=NAVY)
    save(fig, 21, "nightingale_ppt")


def chart22():
    rows = data_rows(BACK, "22 仪表盘图")[2:4]
    value = numeric(rows[0][7], 76) if rows and rows[0][7] is not None else 76
    fig, ax = plt.subplots(figsize=(7, 4.5), subplot_kw=dict(aspect="equal"))
    ax.add_patch(Wedge((0, 0), 1.0, 0, 180, facecolor=LIGHT, edgecolor="white"))
    ax.add_patch(Wedge((0, 0), 1.0, 0, 180 * value / 100, facecolor=TEAL, edgecolor="white"))
    ax.plot([0, math.cos(math.radians(180 - 180 * value / 100))],
            [0, math.sin(math.radians(180 - 180 * value / 100))], color=RED, linewidth=3)
    ax.text(0, -0.1, f"{value:.0f}", ha="center", va="center", fontsize=28, fontweight="bold", color=NAVY)
    ax.set_xlim(-1.15, 1.15)
    ax.set_ylim(-0.2, 1.15)
    ax.axis("off")
    ax.set_title("仪表盘图", loc="left", fontweight="bold", color=NAVY)
    save(fig, 22, "gauge")


def chart23():
    rows = data_rows(BACK, "23 柱形折线图")[2:8]
    cats, bars, line = [clean(r[1]) for r in rows], [numeric(r[2]) for r in rows], [numeric(r[3]) for r in rows]
    x = np.arange(len(cats))
    fig, ax = plt.subplots(figsize=(9, 5))
    ax.bar(x, bars, color=BLUE, width=0.6, label="销售额")
    ax2 = ax.twinx()
    ax2.plot(x, line, color=ORANGE, marker="o", linewidth=2, label="同比")
    ax.set_xticks(x, cats)
    ax.set_ylabel("销售额")
    ax2.set_ylabel("同比")
    style(ax, "柱形折线图")
    ax2.spines["top"].set_visible(False)
    save(fig, 23, "bar_line")


def chart24():
    rows = data_rows(BACK, "24 目标柱形图")[2:9]
    cats, actual, target = [clean(r[1]) for r in rows], [numeric(r[2]) for r in rows], [numeric(r[3]) for r in rows]
    x = np.arange(len(cats))
    fig, ax = plt.subplots(figsize=(9, 5))
    ax.bar(x, target, color=LIGHT, width=0.65, label="目标")
    ax.bar(x, actual, color=TEAL, width=0.42, label="实际")
    ax.set_xticks(x, cats)
    ax.legend(frameon=False, ncol=2)
    style(ax, "目标柱形图")
    save(fig, 24, "target_bar")


def chart25():
    rows = data_rows(BACK, "25 子弹图")[3:10]
    cats, actual, target = [clean(r[1]) for r in rows], [numeric(r[2]) for r in rows], [numeric(r[3]) for r in rows]
    y = np.arange(len(cats))
    fig, ax = plt.subplots(figsize=(9, 5))
    ax.barh(y, target, color=LIGHT, height=0.55)
    ax.barh(y, actual, color=BLUE, height=0.25)
    ax.scatter(target, y, color=ORANGE, s=45, zorder=3, label="目标")
    ax.set_yticks(y, cats)
    ax.legend(frameon=False)
    style(ax, "子弹图")
    ax.grid(axis="x", color=LIGHT)
    ax.grid(axis="y", visible=False)
    save(fig, 25, "bullet")


def chart26():
    rows = data_rows(BACK, "26 柱形圆")[2:8]
    cats, vals, share = [clean(r[1]) for r in rows], [numeric(r[2]) for r in rows], [numeric(r[3]) for r in rows]
    fig, ax = plt.subplots(figsize=(8, 5))
    bars = ax.bar(cats, vals, color=BLUE, width=0.6)
    for b, s in zip(bars, share):
        ax.text(b.get_x() + b.get_width()/2, b.get_height(), f"{s:.0f}", ha="center", va="bottom", fontsize=9)
    style(ax, "柱形圆")
    save(fig, 26, "bar_circle")


def chart27():
    rows = data_rows(BACK, "27 簇状柱形折线图")[2:8]
    cats = [clean(r[1]) for r in rows]
    a, b, line = [numeric(r[2]) for r in rows], [numeric(r[3]) for r in rows], [numeric(r[4]) for r in rows]
    x = np.arange(len(cats))
    fig, ax = plt.subplots(figsize=(9, 5))
    ax.bar(x - 0.18, a, width=0.36, color=BLUE, label="2022")
    ax.bar(x + 0.18, b, width=0.36, color=ORANGE, label="2021")
    ax2 = ax.twinx()
    ax2.plot(x, line, color=PURPLE, marker="o", linewidth=2, label="同比")
    ax.set_xticks(x, cats)
    style(ax, "簇状柱形折线图")
    ax2.spines["top"].set_visible(False)
    save(fig, 27, "clustered_bar_line")


def chart28():
    rows = data_rows(BACK, "28 复合柱形图")[2:17]
    cats, vals, total = [clean(r[1]) for r in rows], [numeric(r[2]) for r in rows], [numeric(r[3]) for r in rows]
    x = np.arange(len(cats))
    fig, ax = plt.subplots(figsize=(9, 5))
    ax.bar(x, vals, color=TEAL, width=0.6)
    ax.plot(x, total, color=ORANGE, marker="o", linewidth=2)
    ax.set_xticks(x, cats)
    style(ax, "复合柱形图")
    ax.tick_params(axis="x", rotation=35)
    save(fig, 28, "combo")


def chart29():
    rows = data_rows(BACK, "29 滑珠图")[2:7]
    cats, vals, remaining, count = [clean(r[1]) for r in rows], [numeric(r[2]) for r in rows], [numeric(r[3]) for r in rows], [numeric(r[4]) for r in rows]
    y = np.arange(len(cats))
    fig, ax = plt.subplots(figsize=(9, 5))
    ax.barh(y, np.ones(len(y)), color=LIGHT, height=0.16)
    ax.scatter(vals, y, s=80, color=TEAL, zorder=3)
    ax.set_yticks(y, cats)
    ax.set_xlim(0, 1)
    ax.xaxis.set_major_formatter(lambda x, pos: f"{x:.0%}")
    style(ax, "滑珠图")
    ax.grid(axis="x", color=LIGHT)
    ax.grid(axis="y", visible=False)
    save(fig, 29, "dot_slider")


def chart30():
    rows = data_rows(BACK, "30 对比滑珠图")[3:8]
    cats = [clean(r[1]) for r in rows]
    a, b = [numeric(r[2]) for r in rows], [numeric(r[3]) for r in rows]
    y = np.arange(len(cats))
    fig, ax = plt.subplots(figsize=(9, 5))
    ax.hlines(y, 0, 1, color=LIGHT, linewidth=5)
    ax.scatter(a, y - 0.08, color=BLUE, s=75, label="2022")
    ax.scatter(b, y + 0.08, color=ORANGE, s=75, label="2021")
    ax.set_yticks(y, cats)
    ax.set_xlim(0, 1)
    ax.xaxis.set_major_formatter(lambda x, pos: f"{x:.0%}")
    ax.legend(frameon=False, ncol=2)
    style(ax, "对比滑珠图")
    ax.grid(axis="x", color=LIGHT)
    ax.grid(axis="y", visible=False)
    save(fig, 30, "comparison_dot_slider")


def main():
    OUT.mkdir(exist_ok=True)
    charts = [
        chart01, chart02, chart03, chart04, chart05, chart06, chart07, chart08,
        chart09, chart10, chart11, chart12, chart13, chart14, chart15, chart16,
        chart17, chart18, chart19, chart20, chart21, chart22, chart23, chart24,
        chart25, chart26, chart27, chart28, chart29, chart30,
    ]
    for chart in charts:
        chart()
    print(f"Generated {len(charts)} charts in {OUT}")


if __name__ == "__main__":
    main()
